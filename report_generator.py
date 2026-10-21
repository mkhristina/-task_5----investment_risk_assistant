
import os
import pandas as pd

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    PageBreak,
)


REPORTS_DIR = "reports"
DATA_PATH = "data/data_with_fuzzy.csv"
METRICS_PATH = "data/model_comparison.csv"


def prepare_output_dir():
    os.makedirs(REPORTS_DIR, exist_ok=True)


def load_report_data():
    """Загружает данные проекта и метрики обученных моделей."""

    if not os.path.exists(DATA_PATH):
        raise FileNotFoundError(
            f"Не найден файл данных: {DATA_PATH}. "
            "Сначала запустите fuzzy_model.py."
        )

    if not os.path.exists(METRICS_PATH):
        raise FileNotFoundError(
            f"Не найдены метрики моделей: {METRICS_PATH}. "
            "Сначала запустите ml_model.py."
        )

    df = pd.read_csv(DATA_PATH)
    metrics_df = pd.read_csv(METRICS_PATH)

    return df, metrics_df


def build_risk_summary(df):
    """Создаёт сводку по классам риска, если метка доступна."""

    if "risk_class" not in df.columns:
        return pd.DataFrame(
            columns=["Класс риска", "Количество проектов", "Доля, %"]
        )

    counts = df["risk_class"].value_counts().reindex(
        [0, 1, 2], fill_value=0
    )

    labels = {
        0: "Низкий риск",
        1: "Средний риск",
        2: "Высокий риск",
    }

    total = max(len(df), 1)

    return pd.DataFrame([
        {
            "Класс риска": labels[class_id],
            "Количество проектов": int(counts[class_id]),
            "Доля, %": round(counts[class_id] / total * 100, 2),
        }
        for class_id in [0, 1, 2]
    ])


def build_rules_note():
    """Возвращает описание нечёткой системы без выдумывания правил."""

    return pd.DataFrame([
        {
            "Компонент": "Финансовый индекс",
            "Описание": "Оценка финансовых показателей проекта",
        },
        {
            "Компонент": "Рыночный индекс",
            "Описание": "Оценка рыночных факторов риска",
        },
        {
            "Компонент": "Организационный индекс",
            "Описание": "Оценка команды, технологии и соответствия стратегии",
        },
        {
            "Компонент": "Fuzzy risk",
            "Описание": "Итоговый нечёткий показатель риска",
        },
    ])


def style_worksheet(ws):
    """Форматирует заголовки и подбирает ширину столбцов."""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells[:200]:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))

        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            32,
        )


def add_dataframe_to_sheet(wb, title, df):
    """Добавляет DataFrame на отдельный лист Excel."""

    ws = wb.create_sheet(title=title)

    if df.empty:
        ws.append(["Нет данных для отображения"])
    else:
        ws.append([str(col) for col in df.columns])

        for row in df.itertuples(index=False, name=None):
            ws.append([
                value.item() if hasattr(value, "item") else value
                for value in row
            ])

    style_worksheet(ws)
    return ws


def generate_excel(df, rules_df, metrics_df):
    """Создаёт Excel-отчёт с данными, метриками и сводкой риска."""

    prepare_output_dir()

    wb = Workbook()
    ws_data = wb.active
    ws_data.title = "Data"

    ws_data.append([str(col) for col in df.columns])

    for row in df.itertuples(index=False, name=None):
        ws_data.append([
            value.item() if hasattr(value, "item") else value
            for value in row
        ])

    style_worksheet(ws_data)

    add_dataframe_to_sheet(wb, "Metrics", metrics_df)
    add_dataframe_to_sheet(wb, "Risk Summary", build_risk_summary(df))
    add_dataframe_to_sheet(wb, "Fuzzy Components", rules_df)

    output_path = os.path.join(
        REPORTS_DIR,
        "Investment_Report.xlsx",
    )
    wb.save(output_path)

    print(f"Excel-отчёт сохранён: {output_path}")


def dataframe_to_report_table(df, max_rows=None):
    """Преобразует таблицу для PDF и обрезает слишком длинные значения."""

    view = df.copy()

    if max_rows is not None:
        view = view.head(max_rows)

    if view.empty:
        return [["Нет данных"]]

    table_data = [[str(col) for col in view.columns]]

    for row in view.itertuples(index=False, name=None):
        formatted_row = []

        for value in row:
            if pd.isna(value):
                formatted_row.append("")
            elif isinstance(value, (float, int)):
                formatted_row.append(f"{value:.3f}" if isinstance(value, float) else str(value))
            else:
                formatted_row.append(str(value)[:80])

        table_data.append(formatted_row)

    return table_data


def add_pdf_table(story, title, df, styles, max_rows=None):
    """Добавляет таблицу с повторяющимся заголовком и сеткой."""

    story.append(Paragraph(title, styles["Heading2"]))
    story.append(Spacer(1, 0.2 * cm))

    table_data = dataframe_to_report_table(df, max_rows=max_rows)
    column_count = max(len(row) for row in table_data)

    available_width = landscape(A4)[0] - 3 * cm
    col_width = available_width / max(column_count, 1)

    table = Table(
        table_data,
        repeatRows=1,
        colWidths=[col_width] * column_count,
    )

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("LEADING", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
            colors.white,
            colors.HexColor("#F2F6FA"),
        ]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))

    story.append(table)
    story.append(Spacer(1, 0.4 * cm))


def generate_pdf(df, metrics_df, rules_df=None):
    """Создаёт PDF-отчёт с общей информацией и таблицами."""

    prepare_output_dir()

    if rules_df is None:
        rules_df = build_rules_note()

    output_path = os.path.join(
        REPORTS_DIR,
        "Investment_Report.pdf",
    )

    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        rightMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=18,
        leading=22,
        spaceAfter=12,
    ))

    story = []

    story.append(Paragraph(
        "Отчёт по оценке инвестиционных рисков",
        styles["ReportTitle"],
    ))
    story.append(Paragraph(
        "Отчёт сформирован на основе синтетических данных учебного проекта.",
        styles["Normal"],
    ))
    story.append(Spacer(1, 0.5 * cm))

    summary = build_risk_summary(df)
    add_pdf_table(
        story,
        "1. Распределение проектов по классам риска",
        summary,
        styles,
    )

    add_pdf_table(
        story,
        "2. Сравнение ML-моделей",
        metrics_df,
        styles,
    )

    add_pdf_table(
        story,
        "3. Компоненты нечёткой системы",
        rules_df,
        styles,
    )

    add_pdf_table(
        story,
        "4. Пример данных проектов (первые 10 строк)",
        df,
        styles,
        max_rows=10,
    )

    doc.build(story)
    print(f"PDF-отчёт сохранён: {output_path}")


def main():
    try:
        df, metrics_df = load_report_data()
        rules_df = build_rules_note()

        generate_pdf(df, metrics_df, rules_df)
        generate_excel(df, rules_df, metrics_df)

    except (FileNotFoundError, ValueError, KeyError) as error:
        print(f"Ошибка формирования отчёта: {error}")


if __name__ == "__main__":
    main()
